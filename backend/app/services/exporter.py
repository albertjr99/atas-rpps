"""Exportação da ata em DOCX e PDF.

DOCX: quando o RPPS cadastrou uma ata-modelo .docx, o próprio arquivo é usado
como base. O corpo é substituído pelo conteúdo da ata e cabeçalho, rodapé,
brasão, margens e estilos do modelo permanecem intactos. (A Etapa 2 do plano
evolui isso para preencher seção por seção do modelo anotado.)
"""
from __future__ import annotations

import io
from html import escape
from html.parser import HTMLParser

from docx import Document
from docx.shared import Pt

# ── HTML simples → blocos ─────────────────────────────────────────────────────

_BLOCK_TAGS = {"p", "h1", "h2", "h3", "h4", "li", "div"}


class _Blocks(HTMLParser):
    """Converte o HTML do editor em blocos [(tipo, [(texto, negrito, itálico)])]."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.blocks: list[tuple[str, list[tuple[str, bool, bool]]]] = []
        self._kind: str | None = None
        self._runs: list[tuple[str, bool, bool]] = []
        self._bold = 0
        self._italic = 0

    def _flush(self) -> None:
        if self._runs and "".join(t for t, _, _ in self._runs).strip():
            self.blocks.append((self._kind or "p", self._runs))
        self._runs = []
        self._kind = None

    def handle_starttag(self, tag, attrs):
        if tag in _BLOCK_TAGS:
            self._flush()
            self._kind = tag if tag != "div" else "p"
        elif tag in ("strong", "b"):
            self._bold += 1
        elif tag in ("em", "i"):
            self._italic += 1
        elif tag == "br":
            self._runs.append(("\n", False, False))

    def handle_endtag(self, tag):
        if tag in _BLOCK_TAGS:
            self._flush()
        elif tag in ("strong", "b"):
            self._bold = max(0, self._bold - 1)
        elif tag in ("em", "i"):
            self._italic = max(0, self._italic - 1)

    def handle_data(self, data):
        if data.strip() or self._runs:
            self._runs.append((data, self._bold > 0, self._italic > 0))

    def close(self):
        super().close()
        self._flush()


def html_to_blocks(html: str) -> list[tuple[str, list[tuple[str, bool, bool]]]]:
    parser = _Blocks()
    parser.feed(html or "")
    parser.close()
    return parser.blocks


# ── DOCX ──────────────────────────────────────────────────────────────────────

def _clear_body(doc: Document) -> None:
    """Remove o conteúdo do corpo preservando as propriedades de seção (sectPr)."""
    body = doc.element.body
    for child in list(body):
        if not child.tag.endswith("}sectPr"):
            body.remove(child)


def build_docx(meeting: dict, *, model_docx: bytes | None, profile: dict) -> bytes:
    if model_docx:
        doc = Document(io.BytesIO(model_docx))
        _clear_body(doc)
    else:
        doc = Document()
        normal = doc.styles["Normal"].font
        normal.name = profile.get("font_name") or "Calibri"
        normal.size = Pt(profile.get("font_size_pt") or 11)

    blocks = html_to_blocks(meeting.get("ata_content") or "")
    if not blocks:
        blocks = [("p", [("Ata ainda não gerada.", False, False)])]

    for kind, runs in blocks:
        para = doc.add_paragraph()
        if kind == "li":
            para.add_run("• ")
        for text, bold, italic in runs:
            run = para.add_run(text)
            run.bold = bold or kind in ("h1", "h2", "h3", "h4")
            run.italic = italic

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── PDF ───────────────────────────────────────────────────────────────────────

def build_pdf(meeting: dict, *, institution: str = "") -> bytes:
    from reportlab.lib.enums import TA_JUSTIFY
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    buf = io.BytesIO()
    pdf = SimpleDocTemplate(buf, pagesize=A4, leftMargin=2.5 * cm, rightMargin=2 * cm,
                            topMargin=2 * cm, bottomMargin=2 * cm,
                            title=f"Ata nº {meeting.get('number')}")
    styles = getSampleStyleSheet()
    body = ParagraphStyle(name="Body", parent=styles["Normal"], fontSize=11, leading=16, alignment=TA_JUSTIFY)
    heading = ParagraphStyle(name="H", parent=styles["Heading3"], fontSize=12, leading=16)
    story = []
    if institution:
        story.append(Paragraph(escape(institution), styles["Title"]))
    story.append(Paragraph(f"<b>ATA Nº {meeting.get('number')} — {escape(meeting.get('title', ''))}</b>", heading))
    story.append(Spacer(1, 0.4 * cm))

    blocks = html_to_blocks(meeting.get("ata_content") or "") or [("p", [("Ata ainda não gerada.", False, False)])]
    for kind, runs in blocks:
        markup = ""
        for text, bold, italic in runs:
            piece = escape(text).replace("\n", "<br/>")
            if italic:
                piece = f"<i>{piece}</i>"
            if bold:
                piece = f"<b>{piece}</b>"
            markup += piece
        if kind == "li":
            markup = "• " + markup
        story.append(Paragraph(markup, heading if kind in ("h1", "h2", "h3", "h4") else body))
        story.append(Spacer(1, 0.2 * cm))

    pdf.build(story)
    return buf.getvalue()
