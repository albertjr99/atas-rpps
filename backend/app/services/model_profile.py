"""Perfil do modelo de ata de cada RPPS (o "DNA" do documento).

Une a ideia do knowledge_extractor do AGL (ler a ata-modelo e extrair padrão
visual e linguístico) com uma implementação enxuta e sem estado global.

O perfil é usado em dois lugares:
1. Nos prompts de geração, para a IA escrever no estilo do RPPS.
2. Na exportação DOCX, que usa o próprio arquivo-modelo como base
   (cabeçalho, rodapé, brasão e margens preservados).
"""
from __future__ import annotations

import io
import re
from collections import Counter

_TITLE_RE = re.compile(r"^\s*(ITEM\s+\d+|\d+(?:\.\d+)*\s*[\-–—:.)]\s*.+)$", re.IGNORECASE)
_STOP = {
    "para", "com", "sem", "dos", "das", "que", "uma", "nas", "nos", "por", "foi", "aos",
    "pela", "pelo", "como", "mais", "sua", "seu", "este", "esta", "isso", "ser", "sobre",
}


def _emu_to_cm(value) -> float | None:
    try:
        return round(float(value) / 360000.0, 2) if value is not None else None
    except (TypeError, ValueError):
        return None


def recurrent_expressions(text: str, max_items: int = 12) -> list[str]:
    """Expressões de 2 a 4 palavras que se repetem (padrão linguístico do RPPS)."""
    words = [w for w in re.findall(r"[A-Za-zÀ-ÿ]{3,}", (text or "").lower()) if w not in _STOP]
    counts: Counter[str] = Counter()
    for n in (2, 3, 4):
        for i in range(len(words) - n + 1):
            counts[" ".join(words[i:i + n])] += 1
    return [expr for expr, c in counts.most_common(max_items * 3) if c >= 2][:max_items]


def profile_from_docx(data: bytes) -> dict:
    from docx import Document

    doc = Document(io.BytesIO(data))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    text = "\n".join(paragraphs)

    header_lines: list[str] = []
    footer_lines: list[str] = []
    for sec in doc.sections:
        header_lines += [p.text.strip() for p in sec.header.paragraphs if p.text.strip()]
        footer_lines += [p.text.strip() for p in sec.footer.paragraphs if p.text.strip()]

    fonts: Counter[str] = Counter()
    sizes: Counter[float] = Counter()
    titles: list[str] = []
    for p in doc.paragraphs:
        txt = (p.text or "").strip()
        style = getattr(getattr(p, "style", None), "name", "") or ""
        if txt and (("heading" in style.lower()) or _TITLE_RE.match(txt) or (txt.isupper() and len(txt) < 120)):
            titles.append(txt[:160])
        for run in p.runs:
            if run.font is not None and run.font.name:
                fonts[run.font.name] += 1
            if run.font is not None and run.font.size is not None:
                sizes[round(run.font.size.pt, 1)] += 1

    normal = doc.styles["Normal"].font if "Normal" in [s.name for s in doc.styles] else None
    sec0 = doc.sections[0] if doc.sections else None
    return {
        "source": "docx",
        "font_name": (fonts.most_common(1)[0][0] if fonts else (normal.name if normal else None)),
        "font_size_pt": (sizes.most_common(1)[0][0] if sizes else
                         (normal.size.pt if normal is not None and normal.size else None)),
        "margins_cm": {
            "top": _emu_to_cm(sec0.top_margin) if sec0 else None,
            "bottom": _emu_to_cm(sec0.bottom_margin) if sec0 else None,
            "left": _emu_to_cm(sec0.left_margin) if sec0 else None,
            "right": _emu_to_cm(sec0.right_margin) if sec0 else None,
        },
        "header_lines": list(dict.fromkeys(header_lines))[:6],
        "footer_lines": list(dict.fromkeys(footer_lines))[:6],
        "section_titles": list(dict.fromkeys(titles))[:30],
        "opening_example": " ".join(paragraphs[:3])[:1200],
        "closing_example": " ".join(paragraphs[-3:])[:1200] if len(paragraphs) > 3 else "",
        "recurrent_expressions": recurrent_expressions(text),
        "paragraph_count": len(paragraphs),
        "sample_text": text[:6000],
    }


def profile_from_text(text: str) -> dict:
    paragraphs = [p.strip() for p in (text or "").split("\n") if p.strip()]
    return {
        "source": "text",
        "section_titles": [p[:160] for p in paragraphs if _TITLE_RE.match(p) or (p.isupper() and len(p) < 120)][:30],
        "opening_example": " ".join(paragraphs[:3])[:1200],
        "closing_example": " ".join(paragraphs[-3:])[:1200] if len(paragraphs) > 3 else "",
        "recurrent_expressions": recurrent_expressions(text),
        "paragraph_count": len(paragraphs),
        "sample_text": (text or "")[:6000],
    }


def profile_prompt_block(profile: dict, institution: str = "") -> str:
    """Transforma o perfil em instruções de estilo para a IA."""
    if not profile:
        return "Não há ata-modelo cadastrada. Use o padrão formal de atas de comitês de investimentos de RPPS."
    lines = []
    if institution:
        lines.append(f"Instituição: {institution}")
    if profile.get("section_titles"):
        lines.append("Seções e títulos usados na ata-modelo, nesta ordem: " + " | ".join(profile["section_titles"][:20]))
    if profile.get("opening_example"):
        lines.append(f"Exemplo de abertura da ata-modelo: \"{profile['opening_example'][:700]}\"")
    if profile.get("closing_example"):
        lines.append(f"Exemplo de encerramento da ata-modelo: \"{profile['closing_example'][:700]}\"")
    if profile.get("recurrent_expressions"):
        lines.append("Expressões recorrentes do RPPS (use quando couber): " + "; ".join(profile["recurrent_expressions"][:10]))
    if profile.get("sample_text"):
        lines.append("Trecho da ata-modelo para imitar estrutura, tom e vocabulário:\n\"\"\"\n"
                     + profile["sample_text"][:3500] + "\n\"\"\"")
    return "\n".join(lines)
