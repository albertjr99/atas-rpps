import io
import json

import pytest
from docx import Document

from app.core import llm
from app.core.config import settings


def _meeting_payload():
    return {
        "title": "Reunião ordinária de outubro",
        "meeting_type": "Reunião Ordinária",
        "date": "2026-10-15",
        "start_time": "09:00",
        "location": "Sala do Comitê",
        "president": "Maria Souza",
        "secretary": "João Lima",
        "participants": [{"name": "Maria Souza", "role": "Presidente"}, {"name": "Ana", "present": False}],
        "agenda": [{"title": "Cenário econômico"}, {"title": "Aplicações e resgates"}],
    }


@pytest.fixture
def fake_ai(monkeypatch):
    """Simula o provedor de IA, devolvendo respostas conforme a operação."""
    calls = []

    def fake_chat(system, user, *, tenant_id=None, operation="chat", **kwargs):
        calls.append(operation)
        if operation == "topic_summary":
            return json.dumps({"discussao": "O comitê discutiu o cenário.", "deliberacao": "Aprovada a aplicação de R$ 1 milhão.",
                               "encaminhamento": "Executar a aplicação.", "responsavel": "Maria Souza", "prazo": "30 dias"})
        if operation == "generate_ata":
            return "```json\n" + json.dumps({"ata_completa": "<h2>ATA Nº 1</h2><p>Aos quinze dias <strong>do mês</strong>.</p><ul><li>Item & cia</li></ul>",
                                             "resumo_executivo": "Resumo.", "proxima_pauta_sugerida": "<ul><li>Pendências</li></ul>"}) + "\n```"
        if operation == "agent_rewrite":
            return json.dumps({"improved": "Texto revisado formalmente pelo agente.", "rationale": "Mais formal."})
        raise AssertionError(operation)

    monkeypatch.setattr(llm, "chat", fake_chat)
    monkeypatch.setattr("app.api.meetings.transcribe", lambda audio, name, tenant_id=None: "fala transcrita")
    return calls


# ── Autenticação ──────────────────────────────────────────────────────────────

def test_register_login_me(client, make_tenant):
    headers, user = make_tenant("RPPS Alfa")
    assert user["role"] == "admin" and user["institution"] == "RPPS Alfa"
    me = client.get("/api/auth/me", headers=headers).json()
    assert me["email"] == user["email"]
    login = client.post("/api/auth/login", json={"email": user["email"], "password": "senha-segura-1"})
    assert login.status_code == 200 and login.json()["token"]


def test_wrong_password_and_missing_token(client, make_tenant):
    _, user = make_tenant()
    assert client.post("/api/auth/login", json={"email": user["email"], "password": "errada123"}).status_code == 401
    assert client.get("/api/meetings").status_code == 401
    assert client.get("/api/meetings", headers={"Authorization": "Bearer lixo"}).status_code == 401


def test_short_password_rejected(client):
    resp = client.post("/api/auth/register", json={"name": "X", "email": "x@x.org", "password": "123"})
    assert resp.status_code == 400


def test_signup_can_be_disabled(client, monkeypatch):
    monkeypatch.setattr(settings, "ALLOW_SIGNUP", False)
    resp = client.post("/api/auth/register", json={"name": "Y", "email": "y@y.org", "password": "senha-segura-1"})
    assert resp.status_code == 403


def test_superadmin_seeded_from_env_only(client):
    resp = client.post("/api/auth/login", json={"email": "root@plataforma.org", "password": "uma-senha-bem-longa-123"})
    assert resp.status_code == 200 and resp.json()["user"]["role"] == "superadmin"
    # As credenciais antigas fixas no código do AGL não existem mais.
    assert client.post("/api/auth/login", json={"email": "agl@agl.com", "password": "AGL2026"}).status_code == 401


# ── Isolamento entre RPPS ─────────────────────────────────────────────────────

def test_tenant_isolation(client, make_tenant, fake_ai):
    a, _ = make_tenant()
    b, _ = make_tenant()
    meeting = client.post("/api/meetings", json=_meeting_payload(), headers=a).json()
    mid, tid = meeting["id"], meeting["agenda"][0]["id"]
    client.post(f"/api/meetings/{mid}/topics/{tid}/finalize", headers=a)

    assert client.get("/api/meetings", headers=b).json() == []
    assert client.get("/api/deliberations", headers=b).json() == []
    for method, path in [("get", f"/api/meetings/{mid}"), ("patch", f"/api/meetings/{mid}"),
                         ("delete", f"/api/meetings/{mid}"), ("get", f"/api/meetings/{mid}/export/docx"),
                         ("get", f"/api/meetings/{mid}/export/pdf"), ("post", f"/api/meetings/{mid}/generate-ata"),
                         ("get", f"/api/meetings/{mid}/agent/changes")]:
        kwargs = {"json": {"title": "invadido"}} if method == "patch" else {}
        resp = getattr(client, method)(path, headers=b, **kwargs)
        assert resp.status_code == 404, (method, path, resp.status_code)

    delib = client.get("/api/deliberations", headers=a).json()[0]
    assert client.patch(f"/api/deliberations/{delib['id']}", json={"status": "concluido"}, headers=b).status_code == 404
    # A reunião continua intacta para o dono.
    assert client.get(f"/api/meetings/{mid}", headers=a).json()["title"] == "Reunião ordinária de outubro"


def test_meeting_numbers_are_per_tenant(client, make_tenant):
    a, _ = make_tenant()
    b, _ = make_tenant()
    assert client.post("/api/meetings", json=_meeting_payload(), headers=a).json()["number"] == 1
    assert client.post("/api/meetings", json=_meeting_payload(), headers=a).json()["number"] == 2
    assert client.post("/api/meetings", json=_meeting_payload(), headers=b).json()["number"] == 1


# ── Perfis ────────────────────────────────────────────────────────────────────

def test_roles(client, make_tenant):
    admin, _ = make_tenant()
    created = client.post("/api/admin/users", headers=admin, json={
        "name": "Leitor", "email": "leitor@rpps.org", "password": "senha-segura-1", "role": "viewer"})
    assert created.status_code == 200
    viewer = {"Authorization": "Bearer " + client.post("/api/auth/login", json={
        "email": "leitor@rpps.org", "password": "senha-segura-1"}).json()["token"]}
    assert client.get("/api/meetings", headers=viewer).status_code == 200
    assert client.post("/api/meetings", json=_meeting_payload(), headers=viewer).status_code == 403
    assert client.get("/api/admin/users", headers=viewer).status_code == 403
    assert client.post("/api/admin/users", headers=admin, json={
        "name": "X", "email": "x2@rpps.org", "password": "senha-segura-1", "role": "superadmin"}).status_code == 400
    # Alias "user" do frontend antigo vira "editor".
    alias = client.post("/api/admin/users", headers=admin, json={
        "name": "Ed", "email": "ed@rpps.org", "password": "senha-segura-1", "role": "user"}).json()
    assert alias["role"] == "editor"


def test_admin_cannot_delete_user_of_other_tenant(client, make_tenant):
    a, _ = make_tenant()
    b, b_user = make_tenant()
    assert client.delete(f"/api/admin/users/{b_user['id']}", headers=a).status_code == 404
    assert client.get("/api/auth/me", headers=b).status_code == 200


# ── Fluxo da reunião ──────────────────────────────────────────────────────────

def test_full_meeting_flow(client, make_tenant, fake_ai):
    h, _ = make_tenant()
    meeting = client.post("/api/meetings", json=_meeting_payload(), headers=h).json()
    mid, tid = meeting["id"], meeting["agenda"][0]["id"]

    assert client.post(f"/api/meetings/{mid}/start", headers=h).json()["status"] == "in_progress"

    for _ in range(2):  # dois pedaços de áudio no mesmo item
        r = client.post(f"/api/meetings/{mid}/topics/{tid}/transcribe", headers=h,
                        files={"file": ("trecho.webm", b"\x1a\x45audio", "audio/webm")})
        assert r.status_code == 200
    topic = client.get(f"/api/meetings/{mid}", headers=h).json()["agenda"][0]
    assert topic["transcription"] == "fala transcrita fala transcrita"

    for _ in range(2):  # finalizar duas vezes não duplica a deliberação
        data = client.post(f"/api/meetings/{mid}/topics/{tid}/finalize", headers=h).json()
    assert data["deliberacao"].startswith("Aprovada")
    delibs = client.get("/api/deliberations", headers=h).json()
    assert len(delibs) == 1 and delibs[0]["responsible"] == "Maria Souza"
    assert client.get("/api/dashboard/stats", headers=h).json()["pending_deliberations"] == 1

    ata = client.post(f"/api/meetings/{mid}/generate-ata", headers=h).json()
    assert "ATA" in ata["ata_completa"]
    saved = client.get(f"/api/meetings/{mid}", headers=h).json()
    assert saved["status"] == "completed" and saved["executive_summary"] == "Resumo."

    pdf = client.get(f"/api/meetings/{mid}/export/pdf", headers=h)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    docx = client.get(f"/api/meetings/{mid}/export/docx", headers=h)
    text = "\n".join(p.text for p in Document(io.BytesIO(docx.content)).paragraphs)
    assert "Aos quinze dias do mês." in text and "• Item & cia" in text

    run = client.post(f"/api/meetings/{mid}/agent/run", headers=h).json()
    assert run["changes_count"] >= 1
    change = client.get(f"/api/meetings/{mid}/agent/pending", headers=h).json()[0]
    ok = client.post(f"/api/meetings/{mid}/agent/feedback", headers=h,
                     json={"change_id": change["id"], "accepted": True})
    assert ok.json()["status"] == "accepted"
    assert client.post(f"/api/meetings/{mid}/agent/feedback", headers=h,
                       json={"change_id": change["id"], "accepted": False}).status_code == 409


def test_ai_failure_does_not_corrupt_meeting(client, make_tenant, monkeypatch):
    h, _ = make_tenant()
    meeting = client.post("/api/meetings", json=_meeting_payload(), headers=h).json()

    def boom(*a, **k):
        raise llm.AIError("provedor fora do ar")

    monkeypatch.setattr(llm, "chat", boom)
    resp = client.post(f"/api/meetings/{meeting['id']}/generate-ata", headers=h)
    assert resp.status_code == 502
    assert client.get(f"/api/meetings/{meeting['id']}", headers=h).json()["ata_content"] == ""


def test_no_ai_configured_returns_clear_error(client, make_tenant):
    h, _ = make_tenant()
    meeting = client.post("/api/meetings", json=_meeting_payload(), headers=h).json()
    resp = client.post(f"/api/meetings/{meeting['id']}/generate-ata", headers=h)
    assert resp.status_code == 502 and "OPENAI_API_KEY" in resp.json()["detail"]


# ── Ata-modelo ────────────────────────────────────────────────────────────────

def _model_docx() -> bytes:
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "INSTITUTO DE PREVIDÊNCIA MUNICIPAL - BRASÃO"
    doc.sections[0].footer.paragraphs[0].text = "Rua Central, 100 - Centro"
    doc.add_paragraph("ATA DA REUNIÃO ORDINÁRIA DO COMITÊ DE INVESTIMENTOS")
    doc.add_paragraph("1 - ABERTURA")
    doc.add_paragraph("Aos dez dias do mês de setembro reuniu-se o Comitê de Investimentos.")
    doc.add_paragraph("2 - DELIBERAÇÕES")
    doc.add_paragraph("Nada mais havendo a tratar, lavrou-se a presente ata.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_model_upload_profile_and_docx_export_preserves_header(client, make_tenant, fake_ai):
    h, _ = make_tenant()
    up = client.post("/api/kb/upload", headers=h, files={"file": ("modelo.docx", _model_docx(),
                     "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert up.status_code == 200 and up.json()["is_model"] is True
    profile = client.get("/api/kb/model", headers=h).json()["profile"]
    assert "1 - ABERTURA" in profile["section_titles"]
    assert profile["header_lines"] == ["INSTITUTO DE PREVIDÊNCIA MUNICIPAL - BRASÃO"]

    meeting = client.post("/api/meetings", json=_meeting_payload(), headers=h).json()
    client.post(f"/api/meetings/{meeting['id']}/generate-ata", headers=h)
    exported = Document(io.BytesIO(client.get(f"/api/meetings/{meeting['id']}/export/docx", headers=h).content))
    assert exported.sections[0].header.paragraphs[0].text == "INSTITUTO DE PREVIDÊNCIA MUNICIPAL - BRASÃO"
    assert exported.sections[0].footer.paragraphs[0].text == "Rua Central, 100 - Centro"
    body = "\n".join(p.text for p in exported.paragraphs)
    assert "Aos dez dias do mês de setembro" not in body  # corpo do modelo foi substituído
    assert "Aos quinze dias do mês." in body


def test_kb_rejects_unknown_types_and_isolates(client, make_tenant):
    a, _ = make_tenant()
    b, _ = make_tenant()
    assert client.post("/api/kb/upload", headers=a, files={"file": ("x.exe", b"MZ", "application/octet-stream")}).status_code == 400
    doc = client.post("/api/kb/upload", headers=a, files={"file": ("nota.txt", b"texto da ata anterior", "text/plain")}).json()
    assert client.get(f"/api/kb/{doc['id']}", headers=b).status_code == 404
    assert client.delete(f"/api/kb/{doc['id']}", headers=b).status_code == 404
    assert client.get("/api/kb", headers=b).json() == []


# ── Diversos ──────────────────────────────────────────────────────────────────

def test_spa_and_api_404(client):
    assert client.get("/").text == "<html>spa</html>"
    assert client.get("/meetings/abc").text == "<html>spa</html>"
    assert client.get("/api/nao-existe").status_code == 404
    traversal = client.get("/%2e%2e/%2e%2e/app/core/config.py")
    assert "SECRET_KEY" not in traversal.text


def test_scenario_without_data_is_marked_unavailable():
    from app.services.economic_data import generate_scenario_analysis
    assert generate_scenario_analysis(None, None)["scenario"] == "indisponivel"
    ok = generate_scenario_analysis(15.0, 5.0)
    assert ok["real_rate"] == 10.0 and "disclaimer" in ok


def test_parse_json_variants():
    assert llm.parse_json('```json\n{"a": 1}\n```') == {"a": 1}
    assert llm.parse_json('Claro! {"a": 2} pronto') == {"a": 2}
    with pytest.raises(llm.AIError):
        llm.parse_json("sem json")
